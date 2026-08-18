"""Gate-the-gater + escape-replay clause tests for the deterministic paper gate.

``final_gate.py`` is the doneness oracle the whole termination argument rests on, so its
aggregation logic is itself tested: a failing clause MUST raise F>0, and a passing stub
MUST NOT mask a failing one. The v2 clause tests replay the actual historical escapes at
clause level against planted fixtures (red-on-defect / green-on-clean). Pure unit tests
(clause functions with injected fixture paths; no subprocesses), so they run in CI
without the paper toolchain. This file is EXCLUDED from the gate's own G5 subset (it
would recurse); it runs in the full unit suite and CI instead.
"""

from __future__ import annotations

import pytest
from paper import gate_lib
from paper.final_gate import (
    CLAUSES,
    FAST_CLAUSES,
    REPO,
    ClauseResult,
    _evaluate,
    _has_value,
    f_count,
    g0_sticky_flag,
    g2_verify_paper,
    g8_derivations,
    g9_surfaces,
    g10_concepts,
    g11_checksums,
    g12_docx,
    g13_claims,
    g14_source_data,
    run_clauses,
)


def _clause(name: str, ok: bool):
    return lambda: ClauseResult(name, ok, "stub")


def test_all_pass_gives_f_zero():
    results = run_clauses((_clause("A", True), _clause("B", True)))
    assert f_count(results) == 0


def test_one_failing_clause_gives_f_positive():
    results = run_clauses((_clause("A", True), _clause("B", False)))
    assert f_count(results) == 1


def test_passing_clause_does_not_mask_failing_one():
    # The core gate-the-gater property: a green clause cannot hide a red one.
    results = run_clauses((_clause("green", True), _clause("red", False), _clause("green2", True)))
    assert f_count(results) == 1
    assert [r.name for r in results if not r.ok] == ["red"]


def test_all_failing_counts_every_clause():
    results = run_clauses((_clause("A", False), _clause("B", False), _clause("C", False)))
    assert f_count(results) == 3


def test_clauseresult_carries_fields():
    r = ClauseResult("G1 paper_numbers", False, "detail")
    assert r.name == "G1 paper_numbers"
    assert r.ok is False
    assert r.summary == "detail"


# ---------------------------------------------------------------------------
# v2 structure invariants
# ---------------------------------------------------------------------------


def test_fast_subset_is_a_subset_of_the_full_gate():
    assert set(FAST_CLAUSES) <= set(CLAUSES)
    names = {c.__name__ for c in FAST_CLAUSES}
    assert names == {
        "g0_sticky_flag",
        "g1_paper_numbers",
        "g3_number_audit",
        "g6_archive_unchanged",
        "g8_derivations",
    }


# ---------------------------------------------------------------------------
# Calibrated-red evaluation (staged CI): exit 0 iff failing set == expected set EXACTLY
# ---------------------------------------------------------------------------


def _results(**named_ok: bool) -> list[ClauseResult]:
    return [ClauseResult(f"{k} x", ok, "stub") for k, ok in named_ok.items()]


def test_calibrated_red_match_exits_zero(capsys):
    rc = _evaluate(_results(G1=True, G8=False, G12=False), expected={"G8", "G12"})
    assert rc == 0
    assert "CALIBRATED-RED match" in capsys.readouterr().out


def test_calibrated_red_unexpected_failure_exits_nonzero(capsys):
    rc = _evaluate(_results(G1=False, G8=False), expected={"G8"})
    assert rc == 1
    assert "MISMATCH" in capsys.readouterr().out


def test_calibrated_red_missing_expected_failure_exits_nonzero():
    # An expected-red clause that unexpectedly PASSES is also a mismatch: the
    # expectation file is stale and must be updated — never silently tolerated.
    rc = _evaluate(_results(G8=True, G12=False), expected={"G8", "G12"})
    assert rc == 1


def test_calibrated_red_intersects_with_fast_subset():
    # --fast runs only a clause subset; expectations outside it must not fail the run.
    rc = _evaluate(_results(G3=False, G8=False), expected={"G2", "G3", "G8", "G12"})
    assert rc == 0


def test_true_green_prints_head_stamped_gate_line(capsys):
    rc = _evaluate(_results(G1=True, G2=True), expected=None)
    assert rc == 0
    out = capsys.readouterr().out
    assert "GATE GREEN sha=" in out and "dirty=" in out


def test_true_green_mode_fails_on_any_red():
    assert _evaluate(_results(G1=True, G2=False), expected=None) == 1


# ---------------------------------------------------------------------------
# Escape replay: G0 sticky flag
# ---------------------------------------------------------------------------


def test_g0_red_while_flag_present_green_when_absent(tmp_path):
    flag = tmp_path / "GATE_RED.flag"
    assert g0_sticky_flag(flag=flag).ok is True
    flag.write_text("written=now blocked_attempts=8\n", encoding="utf-8")
    res = g0_sticky_flag(flag=flag)
    assert res.ok is False and "clear-red-flag" in res.summary


# ---------------------------------------------------------------------------
# Escape replay: G2 fail-open (the v1 bug — absent manifest skipped the comparison)
# ---------------------------------------------------------------------------


def test_g2_fails_closed_when_manifest_absent(tmp_path):
    res = g2_verify_paper(
        allowed_warnings=tmp_path / "absent.txt", _result=(0, "✓ passed, no warnings\n")
    )
    assert res.ok is False and "ABSENT" in res.summary


def test_g2_exact_multiset_match_with_counts(tmp_path):
    aw = tmp_path / "allowed_warnings.txt"
    aw.write_text("# manifest\nwarning alpha\nwarning beta\n", encoding="utf-8")
    out_ok = "⚠  2 warning(s):\n  - warning   alpha\n  - warning beta\n✓ passed\n"
    assert g2_verify_paper(allowed_warnings=aw, _result=(0, out_ok)).ok is True
    # an EXTRA emitted warning fails (v1's substring rule could bless whole families)
    out_extra = out_ok.replace("✓ passed", "  - warning beta\n✓ passed")
    assert g2_verify_paper(allowed_warnings=aw, _result=(0, out_extra)).ok is False
    # a manifested-but-absent warning ALSO fails (stale manifest must be pruned)
    out_missing = "⚠  1 warning(s):\n  - warning alpha\n✓ passed\n"
    assert g2_verify_paper(allowed_warnings=aw, _result=(0, out_missing)).ok is False
    # hard errors fail regardless of the manifest
    assert g2_verify_paper(allowed_warnings=aw, _result=(1, out_ok)).ok is False


# ---------------------------------------------------------------------------
# Escape replay: G8 stale derived artifact (de novo M1 — gec.json vs re-frozen equity)
# ---------------------------------------------------------------------------


def _write_derivations(tmp_path, sha: str) -> tuple[object, object]:
    frozen = tmp_path / "results" / "paper_frozen"
    frozen.mkdir(parents=True)
    equity = frozen / "equity_analysis.json"
    equity.write_text('{"high": {"recall": 0.5023}}', encoding="utf-8")
    gec = frozen / "group_error_calibration.json"
    gec.write_text('{"fnr_gap": 0.0623}', encoding="utf-8")
    d = tmp_path / "DERIVATIONS.tsv"
    d.write_text(
        "results/paper_frozen/group_error_calibration.json\t"
        "paper/compute_group_error_calibration.py\t"
        f"results/paper_frozen/equity_analysis.json\t{sha}\trepo\tG8\n",
        encoding="utf-8",
    )
    return d, equity


def test_g8_matching_stamp_is_green(tmp_path):
    d, equity = _write_derivations(tmp_path, "PENDING")
    good = gate_lib.sha256_file(equity)
    d.write_text(d.read_text(encoding="utf-8").replace("PENDING", good), encoding="utf-8")
    assert g8_derivations(derivations=d, root=tmp_path).ok is True


def test_g8_stale_stamp_is_red(tmp_path):
    """The M1 replay: the input was re-frozen AFTER the derivation was stamped."""
    d, equity = _write_derivations(tmp_path, "PENDING")
    old = gate_lib.sha256_file(equity)
    d.write_text(d.read_text(encoding="utf-8").replace("PENDING", old), encoding="utf-8")
    # the Minnesota merge re-freezes the source out from under the derived artifact
    equity.write_text('{"high": {"recall": 0.4027}}', encoding="utf-8")
    res = g8_derivations(derivations=d, root=tmp_path)
    assert res.ok is False and "STALE" in res.summary.upper()


def test_g8_pending_stamp_fails_closed(tmp_path):
    d, _ = _write_derivations(tmp_path, "PENDING")
    res = g8_derivations(derivations=d, root=tmp_path)
    assert res.ok is False and "PENDING" in res.summary


# ---------------------------------------------------------------------------
# Escape replay: G9 / G10 concept conflicts
# ---------------------------------------------------------------------------


def _registry(tmp_path, status: str):
    reg = tmp_path / "CONCEPT_REGISTRY.tsv"
    reg.write_text(
        "systems_total\tfrozen_snapshot\tsnap.json\t95,223\t"
        r"\b10,746 systems\b;\b90,108\b" + f"\t{status}\tnote\n",
        encoding="utf-8",
    )
    return reg


def test_g9_flags_conflicting_rendering_on_changelog_latest(tmp_path):
    reg = _registry(tmp_path, "ACTIVE")
    (tmp_path / "CHANGELOG.md").write_text(
        "# Changelog\n\n## [3.0.0] - 2026-03-03\n- expanded dataset to 10,746 systems\n",
        encoding="utf-8",
    )
    res = g9_surfaces(root=tmp_path, registry=reg)
    assert res.ok is False and "10,746" in res.summary


def test_g9_ignores_older_changelog_releases(tmp_path):
    reg = _registry(tmp_path, "ACTIVE")
    (tmp_path / "CHANGELOG.md").write_text(
        "# Changelog\n\n## [4.0.0] - 2026-07-01\n- dataset now 95,223 systems\n\n"
        "## [3.0.0] - 2026-03-03\n- expanded dataset to 10,746 systems\n",
        encoding="utf-8",
    )
    assert g9_surfaces(root=tmp_path, registry=reg).ok is True


def test_g10_fails_closed_while_rows_pending(tmp_path):
    reg = _registry(tmp_path, "PENDING_DECIDE")
    res = g10_concepts(root=tmp_path, registry=reg)
    assert res.ok is False and "PENDING_DECIDE" in res.summary


def test_g10_active_scans_manuscript_surfaces(tmp_path):
    reg = _registry(tmp_path, "ACTIVE")
    paper_dir = tmp_path / "paper"
    paper_dir.mkdir()
    (paper_dir / "skeleton.md").write_text("we studied 90,108 systems", encoding="utf-8")
    assert g10_concepts(root=tmp_path, registry=reg).ok is False
    (paper_dir / "skeleton.md").write_text("we studied 95,223 systems", encoding="utf-8")
    assert g10_concepts(root=tmp_path, registry=reg).ok is True


# ---------------------------------------------------------------------------
# Escape replay: G11 frozen checksums (the paper's own documented reviewer command)
# ---------------------------------------------------------------------------


def test_g11_verifies_both_directions(tmp_path):
    frozen = tmp_path / "frozen"
    frozen.mkdir()
    a = frozen / "a.json"
    a.write_text('{"x": 1}', encoding="utf-8")
    cs = frozen / "checksums.sha256"
    cs.write_text(f"{gate_lib.sha256_file(a)}  a.json\n", encoding="utf-8")
    assert g11_checksums(frozen=frozen).ok is True
    # tamper -> mismatch
    a.write_text('{"x": 2}', encoding="utf-8")
    assert g11_checksums(frozen=frozen).ok is False
    # restore + add an UNLISTED frozen JSON -> red (coverage is bidirectional)
    a.write_text('{"x": 1}', encoding="utf-8")
    (frozen / "b.json").write_text("{}", encoding="utf-8")
    res = g11_checksums(frozen=frozen)
    assert res.ok is False and "not in checksums" in res.summary


# ---------------------------------------------------------------------------
# Escape replay: G12 stale DOCX (v1 gated the .md while a stale .docx would upload)
# ---------------------------------------------------------------------------


def _docx_fixture(tmp_path, embed: bool, stale: bool):
    docx = pytest.importorskip("docx")
    (tmp_path / "paper").mkdir()
    src = tmp_path / "paper" / "skeleton.md"
    src.write_text("manuscript v1", encoding="utf-8")
    d = tmp_path / "DERIVATIONS.tsv"
    d.write_text(
        "paper/manuscript.docx\tpaper/convert_to_docx.py\tpaper/skeleton.md\t-\trepo\tG12\n",
        encoding="utf-8",
    )
    rows = gate_lib.parse_derivations(d)
    doc = docx.Document()
    doc.add_paragraph("body")
    if embed:
        doc.core_properties.comments = gate_lib.provenance_string(
            gate_lib.manifest_hash("paper/manuscript.docx", rows, tmp_path),
            "deadbeef",
            "2026-07-01T00:00:00Z",
        )
    doc.save(str(tmp_path / "paper" / "manuscript.docx"))
    if stale:
        src.write_text("manuscript v2 — sources changed after generation", encoding="utf-8")
    return d


def test_g12_fresh_embed_is_green(tmp_path):
    d = _docx_fixture(tmp_path, embed=True, stale=False)
    assert g12_docx(derivations=d, root=tmp_path).ok is True


def test_g12_missing_embed_is_red(tmp_path):
    d = _docx_fixture(tmp_path, embed=False, stale=False)
    res = g12_docx(derivations=d, root=tmp_path)
    assert res.ok is False and "no embedded provenance" in res.summary


def test_g12_stale_sources_are_red(tmp_path):
    d = _docx_fixture(tmp_path, embed=True, stale=True)
    res = g12_docx(derivations=d, root=tmp_path)
    assert res.ok is False and "STALE" in res.summary


def test_g10_concept_conflict_is_case_insensitive(tmp_path):
    r"""Escape replay (Phase 4 panel): 'Sixteen model families' in Methods escaped the
    case-sensitive '\bsixteen' alias; concept-conflict scanning must be case-insensitive."""
    reg = tmp_path / "CONCEPT_REGISTRY.tsv"
    reg.write_text(
        "model_family_count\tcode_census\tsrc/\t20 families\t"
        r"\bsixteen model famil"
        "\tACTIVE\tnote\n",
        encoding="utf-8",
    )
    paper_dir = tmp_path / "paper"
    paper_dir.mkdir()
    (paper_dir / "skeleton.md").write_text(
        "Sixteen model families are evaluated.", encoding="utf-8"
    )
    assert g10_concepts(root=tmp_path, registry=reg).ok is False
    (paper_dir / "skeleton.md").write_text(
        "Twenty model families are evaluated.", encoding="utf-8"
    )
    assert g10_concepts(root=tmp_path, registry=reg).ok is True


def test_manifest_hash_is_line_ending_independent(tmp_path):
    """Escape replay (CI cross-platform): a DOCX generated on Windows (CRLF sources) must
    verify against the same sources checked out on Linux (LF). manifest_hash normalizes text
    line endings; binary sources stay raw."""
    (tmp_path / "paper").mkdir()
    src = tmp_path / "paper" / "skeleton.md"
    d = tmp_path / "DERIVATIONS.tsv"
    d.write_text(
        "paper/manuscript.docx\tpaper/convert_to_docx.py\tpaper/skeleton.md\t-\trepo\tG12\n",
        encoding="utf-8",
    )
    rows = gate_lib.parse_derivations(d)
    src.write_bytes(b"line one\r\nline two\r\n")
    h_crlf = gate_lib.manifest_hash("paper/manuscript.docx", rows, tmp_path)
    src.write_bytes(b"line one\nline two\n")
    h_lf = gate_lib.manifest_hash("paper/manuscript.docx", rows, tmp_path)
    assert h_crlf == h_lf, "text manifest hash must be line-ending independent"
    # A real content change must still change the hash.
    src.write_bytes(b"line one\nline three\n")
    assert gate_lib.manifest_hash("paper/manuscript.docx", rows, tmp_path) != h_lf


# ---------------------------------------------------------------------------
# G13 claims-layer gate (v3) — escape-replay.
# A pinned load-bearing sentence must be present (markers stripped, whitespace
# collapsed); mutating it goes red, updating row+sentence together goes green.
# ---------------------------------------------------------------------------


def _g13_fixture(
    tmp_path,
    sentence_in_file: str,
    pinned: str,
    strength="DESCRIPTIVE",
    evidence="e.json",
    make_evidence=True,
    decisions="# d\n",
):
    """Build a hermetic (root, claims, frozen, decisions) fixture for g13_claims."""
    (tmp_path / "paper").mkdir(exist_ok=True)
    (tmp_path / "paper" / "skeleton.md").write_text(sentence_in_file, encoding="utf-8")
    fz = tmp_path / "results" / "paper_frozen"
    fz.mkdir(parents=True, exist_ok=True)
    if make_evidence:
        (fz / evidence).write_text("{}", encoding="utf-8")
    dec = tmp_path / "paper" / "DECISIONS.md"
    dec.write_text(decisions, encoding="utf-8")
    claims = tmp_path / "paper" / "CLAIMS.tsv"
    claims.write_text(
        "# id\tfile\tsentence\tevidence\tstrength\n"
        f"c1\tpaper/skeleton.md\t{pinned}\t{evidence}\t{strength}\n",
        encoding="utf-8",
    )
    return claims, tmp_path, fz, dec


def test_g13_present_sentence_is_green(tmp_path):
    claims, root, fz, dec = _g13_fixture(
        tmp_path,
        "Intro. The model learns who is monitored. Done.",
        "The model learns who is monitored.",
    )
    assert g13_claims(claims=claims, root=root, frozen=fz, decisions=dec).ok is True


def test_g13_mutated_sentence_is_red(tmp_path):
    # The escape: a conclusion silently reworded while its CLAIMS row stayed put.
    claims, root, fz, dec = _g13_fixture(
        tmp_path,
        "Intro. The model learns the environment. Done.",
        "The model learns who is monitored.",
    )
    res = g13_claims(claims=claims, root=root, frozen=fz, decisions=dec)
    assert res.ok is False and "not found" in res.summary


def test_g13_marker_and_whitespace_insensitive(tmp_path):
    # A legitimate <!--pn:id--> re-point and line-wrap must NOT trip G13.
    claims, root, fz, dec = _g13_fixture(
        tmp_path,
        "sampled <!--pn:mon_ratio_poc-->1.85<!--/pn-->×\nmore intensively",  # noqa: RUF001
        "sampled 1.85× more intensively",  # noqa: RUF001
    )
    assert g13_claims(claims=claims, root=root, frozen=fz, decisions=dec).ok is True


def test_g13_significant_missing_evidence_is_red(tmp_path):
    claims, root, fz, dec = _g13_fixture(
        tmp_path,
        "The signal survives.",
        "The signal survives.",
        strength="SIGNIFICANT",
        make_evidence=False,
    )
    res = g13_claims(claims=claims, root=root, frozen=fz, decisions=dec)
    assert res.ok is False and "evidence missing" in res.summary


def test_g13_decided_missing_decisions_id_is_red(tmp_path):
    claims, root, fz, dec = _g13_fixture(
        tmp_path,
        "Framing frozen here.",
        "Framing frozen here.",
        strength="DECIDED",
        evidence="R5-D1",
        make_evidence=False,
        decisions="# no such section\n",
    )
    res = g13_claims(claims=claims, root=root, frozen=fz, decisions=dec)
    assert res.ok is False and "DECISIONS.md" in res.summary


def test_g13_no_rows_is_red_fail_closed(tmp_path):
    (tmp_path / "paper").mkdir()
    claims = tmp_path / "paper" / "CLAIMS.tsv"
    claims.write_text("# only a comment, no rows\n", encoding="utf-8")
    res = g13_claims(claims=claims, root=tmp_path, frozen=tmp_path, decisions=tmp_path / "d.md")
    assert res.ok is False and "no claims" in res.summary


# ---------------------------------------------------------------------------
# Degenerate-cell detector hardening — must flag T3's real
# pattern (five byte-identical models + NaN macro), pass on genuine numbers.
# ---------------------------------------------------------------------------


def test_flag_degenerate_detects_byte_identical_models():
    # The real T3 escape: five models emitting an identical micro tuple.
    mm = {
        "deep_tobit": {"micro_auroc": 0.753, "micro_auprc": 0.321},
        "tabpfn": {"micro_auroc": 0.753, "micro_auprc": 0.321},
        "voting": {"micro_auroc": 0.753, "micro_auprc": 0.321},
        "stacking": {"micro_auroc": 0.753, "micro_auprc": 0.321},
        "icp": {"micro_auroc": 0.753, "micro_auprc": 0.321},
        "lightgbm": {"micro_auroc": 0.766, "micro_auprc": 0.384},
    }
    problems = gate_lib.flag_degenerate_metric_cells(mm, ("micro_auroc", "micro_auprc"))
    assert any("byte-identical" in p for p in problems)
    assert any("deep_tobit" in p and "icp" in p for p in problems)


def test_flag_degenerate_detects_nan():
    mm = {"m1": {"macro_auroc": float("nan")}, "m2": {"macro_auroc": float("nan")}}
    problems = gate_lib.flag_degenerate_metric_cells(mm, ("macro_auroc",))
    assert sum("is NaN" in p for p in problems) == 2


def test_flag_degenerate_passes_on_distinct_numbers():
    mm = {
        "xgboost": {"auroc": 0.864, "auprc": 0.781},
        "lightgbm": {"auroc": 0.870, "auprc": 0.777},
        "random_forest": {"auroc": 0.871, "auprc": 0.775},
    }
    assert gate_lib.flag_degenerate_metric_cells(mm, ("auroc", "auprc")) == []


def test_degenerate_metric_models_returns_offending_set():
    # Rendering-side companion: exactly the byte-identical + NaN models, so the T3
    # leaderboard relabels them "non-converged".
    mm = {
        "deep_tobit": {"micro_auroc": 0.753, "micro_auprc": 0.321},
        "tabpfn": {"micro_auroc": 0.753, "micro_auprc": 0.321},
        "gnn_gcn": {"micro_auroc": 0.7528, "micro_auprc": 0.400},
        "gnn_sage": {"micro_auroc": 0.7528, "micro_auprc": 0.400},
        "xgboost": {"micro_auroc": 0.765, "micro_auprc": 0.445},
        "broken": {"micro_auroc": float("nan"), "micro_auprc": 0.1},
    }
    bad = gate_lib.degenerate_metric_models(mm, ("micro_auroc", "micro_auprc"))
    assert bad == {"deep_tobit", "tabpfn", "gnn_gcn", "gnn_sage", "broken"}
    assert "xgboost" not in bad  # a genuinely distinct converged model is never flagged


# ---------------------------------------------------------------------------
# G14 source-data escape-replays (historical blind spot). These exercise the real
# generate_source_data.py against the committed frozen archive, so they skip when
# the paper toolchain (openpyxl/pandas/click) or the frozen archive is absent —
# matching the test_paper_* convention. The pure sentinel-matching logic is unit
# tested separately (no toolchain needed).
# ---------------------------------------------------------------------------

_FROZEN = REPO / "results" / "paper_frozen"
_SRC = REPO / "paper" / "source_data"


def _g14_fixture(tmp_path):
    """Copy the committed Source Data into a temp root so it can be perturbed."""
    import shutil

    dst = tmp_path / "paper" / "source_data"
    dst.mkdir(parents=True)
    for f in _SRC.glob("*.xlsx"):
        shutil.copy(f, dst / f.name)
    return tmp_path


def _needs_g14():
    pytest.importorskip("openpyxl")
    pytest.importorskip("pandas")
    pytest.importorskip("click")
    if not (_FROZEN / "split_comparison.json").exists() or not any(_SRC.glob("*.xlsx")):
        pytest.skip("frozen archive / committed source_data not present")


def test_g14_clean_copy_is_green(tmp_path):
    _needs_g14()
    root = _g14_fixture(tmp_path)
    assert g14_source_data(root=root, frozen=_FROZEN).ok is True


def test_g14_mutated_cell_is_red(tmp_path):
    _needs_g14()
    import openpyxl

    root = _g14_fixture(tmp_path)
    f = root / "paper" / "source_data" / "Fig2_source_data.xlsx"
    wb = openpyxl.load_workbook(f)
    ws = wb.worksheets[0]
    ws["A1"] = "TAMPERED"
    wb.save(f)
    r = g14_source_data(root=root, frozen=_FROZEN)
    assert r.ok is False and "differs" in r.summary


def test_g14_missing_file_is_red(tmp_path):
    _needs_g14()
    root = _g14_fixture(tmp_path)
    (root / "paper" / "source_data" / "Fig3_source_data.xlsx").unlink()
    r = g14_source_data(root=root, frozen=_FROZEN)
    assert r.ok is False and "file-set mismatch" in r.summary


def test_g14_extra_file_is_red(tmp_path):
    _needs_g14()
    import shutil

    root = _g14_fixture(tmp_path)
    sd = root / "paper" / "source_data"
    shutil.copy(sd / "Fig1_source_data.xlsx", sd / "Fig5_source_data.xlsx")  # the retired orphan
    r = g14_source_data(root=root, frozen=_FROZEN)
    assert r.ok is False and "file-set mismatch" in r.summary


def test_g14_generator_nonzero_is_red(tmp_path):
    _needs_g14()
    root = _g14_fixture(tmp_path)
    empty_frozen = tmp_path / "empty_frozen"
    empty_frozen.mkdir()
    # fail-loud _load_json raises on the missing split_comparison.json -> generator nonzero.
    r = g14_source_data(root=root, frozen=empty_frozen)
    assert r.ok is False and "generate exit=" in r.summary


def test_g14_has_value_sentinel_logic():
    # The per-figure sentinel matcher: present within tolerance -> True; absent -> False.
    assert _has_value([0.1, 1.847053, "x"], 1.847, 1e-3) is True
    assert _has_value([0.1, 0.5, None], 1.847, 1e-3) is False
    assert _has_value(["only", "strings"], 0.744) is False
