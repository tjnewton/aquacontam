"""Tests for paper markdown -> docx conversion (pure python-docx fallback).

These exercise the dependency-free fallback path that builds real Word tables,
applies inline formatting, and single-spaces the review copy. They guard against
regressing to the old behaviour where pipe tables rendered as literal markdown.
"""

from __future__ import annotations

import re
import zipfile
from pathlib import Path

import pytest

pytest.importorskip("docx")

from docx import Document
from paper.convert_to_docx import (
    LINE_SPACING,
    PAPER_DIR,
    REVIEW_LINE_SPACING,
    WIDE_TABLE_COL_THRESHOLD,
    WIDE_TABLE_FONT_SIZE_PT,
    _add_inline_runs,
    _add_markdown_table,
    _find_main_refs_in_paragraph,
    _insert_introduction_heading,
    _insert_main_figures_and_tables,
    _is_table_separator,
    _postprocess_docx,
    _render_markdown_to_doc,
    _resolve_figure_paths,
    _setup_document,
    _strip_leading_headings,
    convert_primer_to_docx,
)

SAMPLE_MD = """# Title

Some **bold** and *italic* and `code` text.

- first bullet
- second **bold** bullet

1. step one
2. step two

| Model | AUROC | AUPRC |
|:------|------:|------:|
| xgboost | 0.86 | 0.73 |
| rf | 0.87 | 0.79 |

Trailing paragraph.
"""


def _build(md: str, spacing: float = REVIEW_LINE_SPACING) -> Document:
    """Render ``md`` into a configured Document via the shared fallback core."""
    doc = Document()
    _setup_document(doc, line_spacing=spacing)
    _render_markdown_to_doc(doc, md)
    return doc


def test_table_is_real_table() -> None:
    table = _build(SAMPLE_MD).tables[0]
    # 3 columns; header + 2 body rows (GFM separator dropped).
    assert (len(table.columns), len(table.rows)) == (3, 3)
    assert table.rows[0].cells[0].text == "Model"
    assert table.rows[0].cells[0].paragraphs[0].runs[0].bold is True
    assert table.rows[1].cells[0].text == "xgboost"


def test_no_pipe_or_markdown_leak(tmp_path: Path) -> None:
    out = tmp_path / "sample.docx"
    _build(SAMPLE_MD).save(str(out))
    with zipfile.ZipFile(out) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    assert "<w:tbl>" in xml
    assert xml.count("|") < 5
    assert "**bold**" not in xml


def test_inline_formatting_runs() -> None:
    runs = _build("Some **bold** and *italic* and `code`.").paragraphs[-1].runs
    assert any(r.bold for r in runs)
    assert any(r.italic for r in runs)


def test_link_renders_text_only() -> None:
    doc = Document()
    _add_inline_runs(doc.add_paragraph(), "see [the paper](http://example.com/x) now")
    text = doc.paragraphs[-1].text
    assert "the paper" in text
    assert "http" not in text


def test_bullet_and_number_lists() -> None:
    styles = [p.style.name for p in _build(SAMPLE_MD).paragraphs]
    assert "List Bullet" in styles
    assert "List Number" in styles


def test_openxml_pagebreak_fence_no_leak() -> None:
    md = 'Para A\n\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n\nPara B\n'
    joined = "".join(p.text for p in _build(md).paragraphs)
    assert "```" not in joined
    assert "w:type" not in joined


def test_yaml_fence_renders_monospace() -> None:
    md = "Cfg:\n\n```yaml\nn_estimators: 500\n```\n\nDone.\n"
    joined = "\n".join(p.text for p in _build(md).paragraphs)
    assert "```" not in joined
    assert "n_estimators: 500" in joined


def test_wide_table_does_not_crash() -> None:
    header = "| " + " | ".join(f"c{i}" for i in range(44)) + " |"
    sep = "|" + "|".join(["---"] * 44) + "|"
    body = "| " + " | ".join(str(i) for i in range(44)) + " |"
    doc = Document()
    _add_markdown_table(doc, [header, sep, body])
    assert len(doc.tables[0].columns) == 44


def test_ragged_rows_padded() -> None:
    doc = Document()
    _add_markdown_table(doc, ["| a | b | c |", "|---|---|---|", "| 1 | 2 |"])
    assert doc.tables[0].rows[1].cells[2].text == ""


def test_is_table_separator() -> None:
    assert _is_table_separator("|:---|---:|:--:|")
    assert _is_table_separator("| --- | --- |")
    assert not _is_table_separator("| Model | AUROC |")


def test_review_spacing_is_single() -> None:
    doc = _build("Hello.")
    assert doc.styles["Normal"].paragraph_format.line_spacing == REVIEW_LINE_SPACING
    assert REVIEW_LINE_SPACING == 1.0


def test_soft_wrapped_prose_merges_into_one_paragraph() -> None:
    # Hard-wrapped source lines (single newlines) reflow into one paragraph.
    md = (
        "First physical line of a wrapped paragraph\n"
        "second physical line continues it\n"
        "and a third line finishes it.\n"
    )
    texts = [p.text for p in _build(md).paragraphs if p.text.strip()]
    assert texts == [
        "First physical line of a wrapped paragraph "
        "second physical line continues it "
        "and a third line finishes it."
    ]


def test_blank_separated_blocks_stay_separate() -> None:
    # The title-block author/affiliation/corresponding lines must not merge.
    md = "*Tyler J. Newton*\n\nIndependent Researcher\n\n*Corresponding author: Tyler J. Newton*\n"
    texts = [p.text for p in _build(md).paragraphs if p.text.strip()]
    assert len(texts) == 3
    assert "Tyler J. Newton" in texts[0]
    assert "Independent Researcher" in texts[1]
    assert "Corresponding author" in texts[2]


def test_hard_wrapped_list_item_continuation_merges() -> None:
    md = (
        "- **T1 (Binary PFAS detection)**: Predict whether a given PFAS compound\n"
        "  (default: PFOS) is detected above the method detection limit at a water\n"
        "  system. Binary classification; primary metric: AUPRC.\n"
        "- **T2 (Concentration regression)**: Predict the concentration.\n"
    )
    bullets = [
        p for p in _build(md).paragraphs if p.style.name == "List Bullet" and p.text.strip()
    ]
    assert len(bullets) == 2
    assert "method detection limit at a water system." in bullets[0].text
    assert bullets[0].text.endswith("primary metric: AUPRC.")
    assert bullets[1].text.endswith("Predict the concentration.")


def test_consecutive_list_items_stay_separate() -> None:
    md = "- alpha\n- beta\n- gamma\n"
    bullets = [
        p for p in _build(md).paragraphs if p.style.name == "List Bullet" and p.text.strip()
    ]
    assert [p.text for p in bullets] == ["alpha", "beta", "gamma"]


def test_numbered_references_blank_separated_with_wrap() -> None:
    md = (
        "1. Andrews, D. Q. Population-wide exposure to PFAS from drinking\n"
        "   water in the United States. *Env. Sci.* 7, 931 (2020).\n\n"
        "2. Hanna-Attisha, M. Elevated blood lead levels in children associated\n"
        "   with the Flint crisis. *Am. J. Public Health* 106, 283 (2016).\n"
    )
    nums = [p for p in _build(md).paragraphs if p.style.name == "List Number" and p.text.strip()]
    assert len(nums) == 2
    assert "drinking water in the United States." in nums[0].text
    assert "Flint crisis." in nums[1].text


def test_escaped_asterisk_in_bold_renders_clean() -> None:
    # ``\*`` inside ``**...**`` must not leak literal ``**`` / ``\`` and must stay bold.
    doc = Document()
    _add_inline_runs(
        doc.add_paragraph(),
        r"**Rosenbaum Gamma\* (tipping point):** For each feature, Gamma\* is small.",
    )
    para = doc.paragraphs[-1]
    assert "**" not in para.text
    assert "\\*" not in para.text
    assert "Gamma* (tipping point):" in para.text
    assert "Gamma* is small." in para.text
    bold_runs = [r for r in para.runs if r.bold]
    assert any("Rosenbaum Gamma*" in r.text for r in bold_runs)


def test_code_span_inside_italic_no_backtick_leak() -> None:
    # The outer ``*...*`` must not swallow the inner code span and leak backticks.
    doc = Document()
    _add_inline_runs(doc.add_paragraph(), "*See `paper/tables/x.csv` for full results.*")
    para = doc.paragraphs[-1]
    assert "`" not in para.text
    assert para.text == "See paper/tables/x.csv for full results."
    assert all(r.italic for r in para.runs if r.text.strip())
    code_runs = [r for r in para.runs if r.font.name == "Consolas"]
    assert any("paper/tables/x.csv" in r.text for r in code_runs)


def test_prose_then_list_flushes_prose_first() -> None:
    md = (
        "AquaContam defines seven benchmark tasks spanning the major water\n"
        "prediction challenges:\n"
        "- first task\n"
        "- second task\n"
    )
    paras = [p for p in _build(md).paragraphs if p.text.strip()]
    assert paras[0].style.name == "Normal"
    assert paras[0].text == (
        "AquaContam defines seven benchmark tasks spanning the major water prediction challenges:"
    )
    assert [p.text for p in paras if p.style.name == "List Bullet"] == [
        "first task",
        "second task",
    ]


# --- Figure/table injection, intro heading, image width, postprocess spacing ---


def test_wrapped_extended_data_table_not_main_match() -> None:
    # "Extended Data" wrapped to the next line must still strip; no bare Table 3.
    para = "…are summarized in Extended Data\nTable 3."
    figs, tbls = _find_main_refs_in_paragraph(para)
    assert figs == set()
    assert tbls == set()


def test_coexisting_qualified_and_bare_table_ref() -> None:
    para = (
        "collapses AUPRC (Fig. 2b, Extended Data\n"
        "Table 1), the provenance-free model reaches 0.545 (Table 3), and"
    )
    _, tbls = _find_main_refs_in_paragraph(para)
    assert tbls == {3}  # genuine bare Table 3 kept; qualified Table 1 stripped
    assert 1 not in tbls


def test_bare_table_ref_single_line_still_matches() -> None:
    _, tbls = _find_main_refs_in_paragraph("see Table 2 for results")
    assert tbls == {2}


def test_insert_strips_table3_heading() -> None:
    text = "## Results\n\nThe provenance-free model is reported (Table 3).\n"
    out = _insert_main_figures_and_tables(text)
    assert "**Table 3**" in out
    assert "## Table 3" not in out  # the table file's own heading is stripped


def test_strip_leading_headings_helper() -> None:
    out = _strip_leading_headings("## Heading\n\n| a | b |\n|---|---|\n| 1 | 2 |")
    assert out.startswith("| a | b |")
    assert "Heading" not in out


def test_insert_introduction_heading_review_only() -> None:
    text = (
        "# Title\n\n## Abstract\n\nThe abstract body.\n\n"
        "The intro paragraph.\n\n## Results\n\nResults.\n"
    )
    out = _insert_introduction_heading(text)
    assert out.count("## Introduction") == 1
    assert (
        out.index("The abstract body.")
        < out.index("## Introduction")
        < out.index("The intro paragraph.")
    )
    # Idempotent: re-running does not add a second heading.
    assert _insert_introduction_heading(out).count("## Introduction") == 1


def test_resolve_figure_paths_adds_clamped_width() -> None:
    out = _resolve_figure_paths("![Fig 1](figures/fig1_dataset_overview.png)")
    m = re.search(r"\{width=([\d.]+)in\}", out)
    assert m is not None
    assert float(m.group(1)) <= 6.5


def test_fallback_image_line_does_not_leak_width_attr() -> None:
    md = "![Cap](figures/fig1_dataset_overview.png){width=6.5in}\n"
    joined = "\n".join(p.text for p in _build(md).paragraphs)
    assert "width=6.5in" not in joined
    assert "{" not in joined


def test_postprocess_line_spacing_param(tmp_path: Path) -> None:
    doc = Document()
    doc.add_paragraph("Body text.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].paragraphs[0].add_run("cell text")
    review = tmp_path / "review.docx"
    doc.save(str(review))
    _postprocess_docx(review, line_spacing=REVIEW_LINE_SPACING)
    d = Document(str(review))
    assert d.paragraphs[0].paragraph_format.line_spacing == REVIEW_LINE_SPACING
    cell_para = d.tables[0].rows[0].cells[0].paragraphs[0]
    assert cell_para.paragraph_format.line_spacing == REVIEW_LINE_SPACING

    doc2 = Document()
    doc2.add_paragraph("Body.")
    sub = tmp_path / "submission.docx"
    doc2.save(str(sub))
    _postprocess_docx(sub)
    assert Document(str(sub)).paragraphs[0].paragraph_format.line_spacing == LINE_SPACING


def test_postprocess_wide_table_font_shrink(tmp_path: Path) -> None:
    from docx.shared import Pt

    doc = Document()
    wide = doc.add_table(rows=1, cols=WIDE_TABLE_COL_THRESHOLD + 1)
    for cell in wide.rows[0].cells:
        cell.paragraphs[0].add_run("x")
    narrow = doc.add_table(rows=1, cols=2)
    for cell in narrow.rows[0].cells:
        cell.paragraphs[0].add_run("y")
    out = tmp_path / "tables.docx"
    doc.save(str(out))
    _postprocess_docx(out)
    d = Document(str(out))
    wide_run = d.tables[0].rows[0].cells[0].paragraphs[0].runs[0]
    narrow_run = d.tables[1].rows[0].cells[0].paragraphs[0].runs[0]
    assert wide_run.font.size == Pt(WIDE_TABLE_FONT_SIZE_PT)
    assert narrow_run.font.size is None  # narrow table left at default


def test_supplementary_table_sequence_check(tmp_path: Path) -> None:
    from paper.verify_paper import check_supplementary_table_sequence

    # Contiguous 1..4 (incl. the "N / Table 3" slash-separator heading variant).
    good = tmp_path / "good.md"
    good.write_text(
        "### Supplementary Table 1: Alpha\n\nbody\n\n"
        "### Supplementary Table 2: Beta\n\nbody\n\n"
        "### Supplementary Table 3 / Table 3: Gamma\n\nbody\n\n"
        "### Supplementary Table 4: Delta\n\nbody\n",
        encoding="utf-8",
    )
    assert check_supplementary_table_sequence(good) == []

    # A gap at 2 must be flagged.
    gapped = tmp_path / "gap.md"
    gapped.write_text(
        "### Supplementary Table 1: A\n\n### Supplementary Table 3: C\n",
        encoding="utf-8",
    )
    warnings = check_supplementary_table_sequence(gapped)
    assert any("Supplementary Table 2 missing" in w for w in warnings)


def test_plain_language_primer_converts_clean(tmp_path: Path) -> None:
    """The standalone reader's primer converts to a clean, structured docx.

    Guards the ``convert_primer_to_docx`` path (and the primer source itself):
    the expected sections render, and no raw markdown (pipes or ``**``) leaks.
    Page count is not asserted here (it depends on the renderer); see the
    primer plan for the line-budget rationale.
    """
    primer_md = PAPER_DIR / "plain_language_primer.md"
    assert primer_md.exists(), "paper/plain_language_primer.md is missing"

    out = tmp_path / "primer.docx"
    convert_primer_to_docx(primer_md, out)
    assert out.exists() and out.stat().st_size > 0

    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    for heading in (
        "Understanding This Paper",
        "What the paper shows",
        "Key terms, translated",
        "Reading the headline numbers",
    ):
        assert heading in text, f"missing section: {heading}"

    # No literal markdown should survive into the rendered document.
    with zipfile.ZipFile(out) as zf:
        xml = zf.read("word/document.xml").decode("utf-8")
    assert "**" not in xml
    assert xml.count("|") < 5
